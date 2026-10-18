"""Three-sheet competition export. Internal amounts are yuan; sheets use wan yuan."""
from __future__ import annotations

import io
import json
from itertools import count

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .database import one, rows
from .models import BidExtract, ItemExtract, LotExtract, ProjectExtract
from .normalization import yuan_to_wan
from .item_hierarchy import analyze_lot

HEADERS = {
    '总表': ['序号','URL','UUID','标段/合同包','项目名称','采购单位','中标供应商','中标金额'],
    '标的表': ['序号','UUID','标段/合同包','产品服务名称','品目','品牌','规格型号','数量','单价','总价','类别'],
    '投标人表': ['序号','UUID','标段/合同包','名称','类别','通过审查','报价（万元）','得分','未中标原因'],
}
OPTIONAL = {'总表': {1}, '标的表': {1,11}, '投标人表': {1,8,9}}


def export_records(dataset_id: str, job_id: str | None = None) -> tuple[list[ProjectExtract], list[list]]:
    notes = [['类型','公告UUID','状态','说明']]
    if job_id:
        job = one('SELECT * FROM jobs WHERE id=? AND dataset_id=?',(job_id,dataset_id))
        if not job:
            raise ValueError('批次不存在或不属于当前数据集')
        if job['mode'] not in {'extract','replay'}:
            raise ValueError('解析预检批次没有抽取结果，不能生成三表')
        records = []
        for doc in rows('SELECT * FROM job_documents WHERE job_id=? ORDER BY notice_uuid',(job_id,)):
            if doc['result_json']:
                records.append(ProjectExtract.model_validate_json(doc['result_json']))
            notes.append(['批次',doc['notice_uuid'],doc['status'],'；'.join(json.loads(doc['errors']))])
        notes.append(['批次',job_id,job['status'],'批次快照；失败公告不在三表中，请同时核对处理说明'])
        return records, notes
    records = [ProjectExtract.model_validate_json(x['result_json']) for x in
               rows('SELECT result_json FROM notices WHERE dataset_id=? ORDER BY notice_uuid',(dataset_id,))]
    found = {r.notice_uuid for r in records}
    # Historical databases predate notice snapshots. Preserve their exportability.
    legacy = {}
    for p in rows('SELECT * FROM projects WHERE dataset_id=? ORDER BY id',(dataset_id,)):
        if p['notice_uuid'] in found:
            continue
        uid = p['notice_uuid'] or f"legacy-{p['id']}"
        result = legacy.setdefault(uid,ProjectExtract(notice_uuid=p['notice_uuid'], project_name=p['name'],project_code=p['project_code'],
            purchaser=p['purchaser'],source_url=p['source_url'],processing_status=p['processing_status'],
            warnings=['历史记录缺少公告快照；请复核原文']))
        lot = LotExtract(lot_name=p['lot_name'],lot_key=p.get('lot_key',''),budget_amount=p['budget_amount'],awarded_amount=p['awarded_amount'])
        for i in rows('SELECT * FROM items WHERE project_id=? ORDER BY id',(p['id'],)):
            lot.items.append(ItemExtract(**{**i,'supplier':i['winner_supplier'],'derived_values':json.loads(i['derived_values']),
                                           'related_sources':json.loads(i.get('related_sources') or '[]')}))
        lot.bids = [BidExtract.model_validate({**b, 'evaluation_scores': json.loads(b.get('evaluation_scores') or '[]'),
                    'related_sources': json.loads(b.get('related_sources') or '[]')}) for b in rows('SELECT * FROM bids WHERE project_id=? ORDER BY id',(p['id'],))]
        result.lots.append(lot)
    return records+list(legacy.values()), notes


def workbook_bytes(records: list[ProjectExtract], notes: list[list] | None = None, include_notes: bool=False,
                   item_view: str='source') -> bytes:
    if item_view not in {'source', 'detail'}:
        raise ValueError('Unknown item view')
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name:workbook.create_sheet(name) for name in HEADERS}
    # Worksheet.max_row scans existing cells; asking it for every row makes full exports quadratic.
    sequence = {name:count(1) for name in HEADERS}
    for name, headers in HEADERS.items():
        sheets[name].append(headers)
    details = list(notes or [['类型','公告UUID','状态','说明']])
    details.append(['导出约定','','','金额单位为万元；原文缺失留空；推算值不替代原文；UUID保持来源前缀，不自行加后缀'])
    details.append(['审查口径','','','明确未通过者标为其他并排除有效共同竞标；未知审查的非中标者保留在投标表，关系分析暂不纳入'])
    for record in records:
        details.extend([['质量',record.notice_uuid,record.processing_status,n] for n in record.warnings])
        if record.processing_status != 'complete' and not record.warnings:
            details.append(['质量',record.notice_uuid,record.processing_status,'结果需要人工复核'])
        if record.no_target_reason:
            details.append(['空结果',record.notice_uuid,record.processing_status,record.no_target_reason])
        for lot in record.package_records():
            winners = list(dict.fromkeys(b.bidder for b in lot.bids if b.is_winner and b.eligible))
            sheets['总表'].append([next(sequence['总表']),record.source_url or None,record.notice_uuid,
                lot.lot_name or None,record.project_name or None,record.purchaser or None,'；'.join(winners) or None,
                yuan_to_wan(lot.awarded_amount,2)])
            excluded = set()
            if item_view == 'detail':
                hierarchy = analyze_lot(lot)
                excluded = set(hierarchy['excluded_indices'])
                for relation in hierarchy['relations']:
                    details.append(['层级', record.notice_uuid, relation['status'],
                        f"{lot.lot_name}: {relation['parent']} -> {', '.join(relation['children'])}; " +
                        ('原文证据通过，明细视图保留子项；摘要仍在原文记录中' if relation['status'] == 'confirmed'
                         else '保留全部原文行：' + '；'.join(relation['reasons']))])
            for index, item in enumerate(lot.items):
                if index in excluded:
                    continue
                # Source category is optional and its official wording is ambiguous; leave blank.
                sheets['标的表'].append([next(sequence['标的表']),record.notice_uuid,lot.lot_name or None,
                    item.product_name or None,item.category or None,item.brand or None,item.model or None,
                    item.quantity,yuan_to_wan(item.unit_price),yuan_to_wan(item.total_price),None])
                if item.derived_values:
                    details.append(['推算',record.notice_uuid,lot.lot_name,json.dumps(item.derived_values,ensure_ascii=False)+'（元，辅助值）'])
            for bid in lot.bids:
                states = (bid.review_status,bid.qualification_status,bid.compliance_status)
                disclosed_pass = bid.review_status=='通过' or bid.qualification_status==bid.compliance_status=='通过'
                status = '未通过' if '未通过' in states else '通过' if disclosed_pass else None
                if status == '未通过' and bid.review_reason:
                    status += '：'+bid.review_reason
                sheets['投标人表'].append([next(sequence['投标人表']),record.notice_uuid,lot.lot_name or None,
                    bid.bidder,bid.role,status,yuan_to_wan(bid.bid_amount,2),
                    round(bid.score,2) if bid.score is not None else None,bid.non_winning_reason or None])
    output_sheets = list(sheets.items())
    if include_notes:
        description = workbook.create_sheet('处理说明')
        for row in details:
            description.append(row)
        output_sheets.append(('处理说明',description))
    for name,sheet in output_sheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    # Spreadsheet formula execution must not originate in external source text.
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('',cell.value)[:32767]
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top',wrap_text=True)
        for cell in sheet[1]:
            optional = cell.column in OPTIONAL.get(name,set())
            cell.fill = PatternFill('solid',fgColor='808080' if optional else '4472C4')
            cell.font = Font(color='FFFFFF',bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24 if column[0].column>1 else 10
    for row in sheets['总表'].iter_rows(min_row=2):
        row[7].number_format='0.00'
    for row in sheets['投标人表'].iter_rows(min_row=2):
        row[6].number_format='0.00'
        row[7].number_format='0.00'
    for row in sheets['标的表'].iter_rows(min_row=2):
        row[8].number_format='0.00########'
        row[9].number_format='0.00########'
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
