from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import stat
import subprocess
import tempfile
import threading
import zipfile
import uuid
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Iterable

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from .config import WORKSPACE, settings
from .normalization import clean_text

SUPPORTED = {'.html','.htm','.pdf','.doc','.docx','.xls','.xlsx','.txt','.csv','.png','.jpg','.jpeg','.webp','.wps','.rtf'}
PARSER_VERSION = 'parse-v3.6'
PARSE_WARNING_PREFIX = '[PARSE_WARNING] '
_GOODS_FIELDS = ('货物名称', '货物品牌', '货物型号', '货物数量', '货物单价')
_ITEM_SEPARATOR = re.compile(r'[；;]')
NOTICE_ID = re.compile(r'^t\d{8}_\d+$')
_PARSE_SLOTS = threading.BoundedSemaphore(settings.parse_concurrency)


class ParseError(ValueError):
    pass


@dataclass
class ArchiveBudget:
    limit: int = settings.max_expanded_bytes
    size: int = 0
    files: int = 0

    def consume(self, size: int) -> None:
        self.size += size
        self.files += 1
        if self.size > self.limit or self.files > settings.max_archive_files:
            raise ParseError('解压累计体积或文件数超出配置限制')


def archive_name(name: str) -> str:
    try:
        raw = name.encode('cp437')
    except UnicodeEncodeError:
        return name
    for encoding in ('utf-8','gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return name


def safe_target(destination: Path, name: str) -> Path:
    parts = PurePosixPath(name.replace('\\','/'))
    if parts.is_absolute() or '..' in parts.parts or any(':' in x for x in parts.parts):
        raise ParseError('压缩包包含非法路径')
    target = destination.joinpath(*parts.parts).resolve()
    if not target.is_relative_to(destination.resolve()):
        raise ParseError('压缩包路径越界')
    return target


def safe_unzip(path: Path, destination: Path, max_total_bytes: int | None = None,
               budget: ArchiveBudget | None = None) -> list[Path]:
    budget = budget or ArchiveBudget(limit=max_total_bytes or settings.max_expanded_bytes)
    extracted = []
    destination.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if stat.S_ISLNK(info.external_attr >> 16):
                raise ParseError('不接受压缩包符号链接')
            target = safe_target(destination,info.filename if info.flag_bits & 0x800 else archive_name(info.filename))
            if info.is_dir():
                target.mkdir(parents=True,exist_ok=True)
                continue
            if info.flag_bits & 1:
                raise ParseError('附件已加密，需要可读取的原始文件')
            budget.consume(info.file_size)
            if shutil.disk_usage(destination).free < info.file_size+128*1024**2:
                raise ParseError('磁盘剩余空间不足，解压已停止')
            if info.file_size > 50*1024**2 and info.file_size/max(info.compress_size,1) > 1000:
                raise ParseError('压缩比异常')
            if target.exists():
                raise ParseError('压缩包包含重名文件')
            target.parent.mkdir(parents=True,exist_ok=True)
            actual = 0
            with archive.open(info) as source, target.open('wb') as output:
                while chunk := source.read(1024**2):
                    actual += len(chunk)
                    if actual > info.file_size:
                        raise ParseError('解压文件超出声明大小')
                    if actual % (16*1024**2) < len(chunk) and shutil.disk_usage(destination).free < len(chunk)+128*1024**2:
                        raise ParseError('磁盘剩余空间不足，解压已停止')
                    output.write(chunk)
            extracted.append(target)
    return extracted


def seven_zip() -> str | None:
    local = WORKSPACE / 'tools/sevenzip/7z.exe'
    return (str(local) if local.is_file() else None) or shutil.which('7zz') or shutil.which('7z') or (str(Path(os.getenv('ProgramFiles','C:/Program Files'))/'7-Zip/7z.exe') if (Path(os.getenv('ProgramFiles','C:/Program Files'))/'7-Zip/7z.exe').is_file() else None)


def office_tool() -> str | None:
    for relative in ('tools/libreoffice/program/soffice.com', 'tools/libreoffice/LibreOffice/program/soffice.com'):
        local = WORKSPACE / relative
        if local.is_file():
            return str(local)
    return shutil.which('libreoffice') or shutil.which('soffice')


def _external_archive(path: Path, destination: Path, budget: ArchiveBudget) -> list[Path]:
    executable = seven_zip()
    if not executable:
        raise ParseError('RAR/7z 解包需要 7z/7zz；请使用已配置工具的部署环境')
    destination.mkdir(parents=True,exist_ok=True)
    result = subprocess.run([executable,'l','-slt','-sccUTF-8','--',str(path)],capture_output=True,timeout=45,stdin=subprocess.DEVNULL)
    if result.returncode:
        raise ParseError('无法列出 RAR/7z 内容，可能已加密或损坏')
    listing = result.stdout.decode('utf-8',errors='replace')
    extracted = []
    for block in re.split(r'\r?\n\r?\n',listing):
        props = dict(line.split(' = ',1) for line in block.splitlines() if ' = ' in line)
        if 'Path' not in props or 'Size' not in props or 'Type' in props or props.get('Folder') == '+':
            continue
        if props.get('Encrypted') == '+':
            raise ParseError('附件已加密')
        target = safe_target(destination,props['Path'])
        size = int(props['Size'])
        budget.consume(size)
        if shutil.disk_usage(destination.parent).free < size+128*1024**2:
            raise ParseError('磁盘剩余空间不足，解压已停止')
        target.parent.mkdir(parents=True,exist_ok=True)
        if target.exists():
            raise ParseError('压缩包包含重名文件')
        with target.open('wb') as output:
            process = subprocess.Popen([executable,'x','-so','-y','-spd','--',str(path),props['Path']],stdout=subprocess.PIPE,
                                       stderr=subprocess.DEVNULL,stdin=subprocess.DEVNULL)
            timer = threading.Timer(90,lambda: process.kill() if process.poll() is None else None)
            timer.daemon = True
            timer.start()
            try:
                written = 0
                while chunk := process.stdout.read(1024**2):
                    written += len(chunk)
                    if written > size:
                        raise ParseError('解包输出超出声明大小')
                    if written % (16*1024**2) < len(chunk) and shutil.disk_usage(destination).free < len(chunk)+128*1024**2:
                        raise ParseError('磁盘剩余空间不足，解压已停止')
                    output.write(chunk)
                if process.wait(timeout=45) or written != size:
                    raise ParseError('RAR/7z 文件解包失败')
            finally:
                timer.cancel()
                if process.poll() is None:
                    process.kill()
                    process.wait()
                process.stdout.close()
        extracted.append(target)
    return extracted


def detect_type(path: Path) -> str:
    with path.open('rb') as stream:
        head = stream.read(4096)
    if head.startswith(b'%PDF-'):
        return '.pdf'
    if head.startswith((b'PK\x03\x04',b'PK\x05\x06',b'PK\x07\x08')):
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
            if 'word/document.xml' in names:
                return '.docx'
            if 'xl/workbook.xml' in names:
                return '.xlsx'
        return '.zip'
    if head.startswith(b'Rar!\x1a\x07'):
        return '.rar'
    if head.startswith(b'7z\xbc\xaf\x27\x1c'):
        return '.7z'
    if head.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return '.xls' if path.suffix.lower() == '.xls' else '.doc'
    if head.startswith(b'\x89PNG'):
        return '.png'
    if head.startswith(b'\xff\xd8\xff'):
        return '.jpg'
    if head.startswith(b'RIFF') and head[8:12] == b'WEBP':
        return '.webp'
    if head.startswith(b'{\\rtf'):
        return '.rtf'
    if re.search(br'<(?:html|head|body|title|link|script|div|meta|table)(?:\s|>)|<!doctype\s+html',head,re.I):
        return '.html'
    if path.suffix.lower() in {'.txt','.csv','.wps','.rtf'}:
        return path.suffix.lower()
    raise ParseError('文件头无法识别，扩展名可能与实际内容不符')


def expand_archives(paths: Iterable[Path], workdir: Path, issues: list[str] | None = None) -> list[Path]:
    queue = [(p,0) for p in paths]
    files = []
    budget = ArchiveBudget()
    index = 0
    while queue:
        path, depth = queue.pop(0)
        try:
            kind = detect_type(path)
            if kind in {'.zip','.rar','.7z'}:
                index += 1
                if depth >= 12 or index > 2000:
                    raise ParseError('附件压缩层级或数量超限')
                destination = workdir / f'archive_{index}'
                children = safe_unzip(path,destination,budget=budget) if kind == '.zip' else _external_archive(path,destination,budget)
                queue.extend((child,depth+1) for child in children)
            else:
                files.append(path)
        except Exception as exc:
            if issues is None:
                raise
            issues.append(f'{path.name}: {type(exc).__name__}: {exc}')
    return files


def decode_text(raw: bytes) -> str:
    for encoding in ('utf-8-sig','gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return raw.decode('utf-8',errors='replace')


def _bounded_join(parts) -> str:
    result = []
    size = 0
    for part in parts:
        size += len(part)+1
        if size > settings.max_document_chars:
            raise ParseError('文档文本量超过配置上限，请拆分处理')
        result.append(part)
    return '\n'.join(result)


def _table_rows(rows: Iterable[tuple[int, list[str]]], table_index: int) -> Iterable[str]:
    """Split only fully aligned goods columns; shared cells keep their text."""
    columns = {}
    width = 0
    for row_index, cells in rows:
        if not any(cell.strip() for cell in cells):
            continue
        location = f'TABLE:{table_index} ROW:{row_index}'
        labels = [re.sub(r'\s+', '', cell) for cell in cells]
        header = {field: [index for index, label in enumerate(labels) if label.startswith(field)]
                  for field in _GOODS_FIELDS}
        if any(header.values()):
            columns = {field: indices[0] for field, indices in header.items() if len(indices) == 1}
            if len(columns) != len(_GOODS_FIELDS) or len(set(columns.values())) != len(_GOODS_FIELDS):
                columns = {}
            width = len(cells)
        elif columns and columns['货物名称'] < len(cells) and _ITEM_SEPARATOR.search(cells[columns['货物名称']]):
            parts = {field: [part.strip() for part in _ITEM_SEPARATOR.split(cells[index])]
                     for field, index in columns.items() if index < len(cells)}
            count = len(parts['货物名称'])
            shared = [cell for index, cell in enumerate(cells) if index not in columns.values()]
            aligned = (len(cells) == width and len(parts) == len(_GOODS_FIELDS)
                       and all(len(values) == count and all(values) for values in parts.values())
                       and not any(_ITEM_SEPARATOR.search(cell) for cell in shared))
            if aligned:
                for item_index in range(count):
                    item = list(cells)
                    for field, index in columns.items():
                        item[index] = parts[field][item_index]
                    yield f'[{location} ITEM:{item_index+1}] ' + ' | '.join(item)
                continue
            yield PARSE_WARNING_PREFIX + f'表{table_index}第{row_index}行：货物分号并列的列数、空值或共享字段无法严格对应，保留原行未拆分，需核验'
        yield f'[{location}] ' + ' | '.join(cells)


def _html_cell_text(cell) -> str:
    """Preserve source spaces and block boundaries; join hyphenated wraps only."""
    parts = []
    stack = list(reversed(list(cell.children)))
    while stack:
        node = stack.pop()
        if isinstance(node, Comment):
            continue
        elif isinstance(node, NavigableString):
            parts.append(str(node))
        elif isinstance(node, str):
            parts.append(node)  # Closing block boundary queued below.
        elif isinstance(node, Tag) and node.name not in {'script', 'style', 'noscript'}:
            if node.name == 'br':
                parts.append('\n')
                continue
            if node.name in {'p', 'div', 'li', 'ul', 'ol', 'section', 'h1', 'h2', 'h3'}:
                parts.append('\n')
                stack.append('\n')
            stack.extend(reversed(list(node.children)))
    text = ''.join(parts)
    text = re.sub(r'(?<=[A-Za-z0-9])[ \t\r]*\n\s*(?=-[A-Za-z0-9])', '', text)
    text = re.sub(r'(?<=[A-Za-z0-9]-)[ \t\r]*\n\s*(?=[A-Za-z0-9])', '', text)
    return re.sub(r'\s+', ' ', text).strip()


def _html(path: Path) -> str:
    soup = BeautifulSoup(decode_text(path.read_bytes()),'lxml')
    if any(word in soup.get_text() for word in ('系统正在维护中','Access Denied','403 Forbidden','404 Not Found')):
        raise ParseError('附件内容是维护/访问错误页面')
    if path.suffix.lower() not in {'.html','.htm'} and not soup.find('table'):
        raise ParseError('附件是网页响应，未发现可用数据表')
    if path.suffix.lower() not in {'.html','.htm'} and re.search(r'请输入验证码|验证码错误|人机验证|验证后下载|登录后下载|访问过于频繁', soup.get_text()):
        raise ParseError('附件内容是下载验证页面，原始文档未提供')
    for tag in soup(['script','style','noscript']):
        tag.decompose()
    # Retain source headers while serializing complete leaf tables only once.
    for index, table in enumerate(list(soup.find_all('table')),1):
        if table.find('table'):
            continue
        rows = []
        pending = {}
        for row_index, tr in enumerate(table.find_all('tr'),1):
            cells = []
            col = 0
            for cell in tr.find_all(['th','td'],recursive=False):
                while (row_index,col) in pending:
                    cells.append(pending[(row_index,col)]); col += 1
                text = _html_cell_text(cell)
                colspan = max(1,min(100,int(cell.get('colspan',1))))
                rowspan = max(1,min(100,int(cell.get('rowspan',1))))
                for _ in range(colspan):
                    cells.append(text)
                    for next_row in range(row_index+1,row_index+rowspan):
                        pending[(next_row,col)] = text
                    col += 1
            while (row_index,col) in pending:
                cells.append(pending[(row_index,col)]); col += 1
            if cells:
                rows.append((row_index, cells))
        chunks = _bounded_join(_table_rows(rows, index))
        table.replace_with(soup.new_string('\n' + chunks + '\n'))
    return _bounded_join([soup.get_text('\n',strip=True)])


def _has_image(resources, seen: set | None = None, depth: int=0) -> bool:
    seen = seen if seen is not None else set()
    if depth>8 or not resources:
        return False
    resources = resources.get_object() if hasattr(resources,'get_object') else resources
    images = resources.get('/XObject',{})
    images = images.get_object() if hasattr(images,'get_object') else images
    for ref in images.values():
        key = (getattr(ref,'idnum',None),getattr(ref,'generation',None)) if hasattr(ref,'idnum') else id(ref)
        if key in seen:
            continue
        seen.add(key)
        obj = ref.get_object() if hasattr(ref,'get_object') else ref
        if obj.get('/Subtype') == '/Image' or (obj.get('/Subtype') == '/Form' and _has_image(obj.get('/Resources',{}),seen,depth+1)):
            return True
    return False


def _pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    if len(reader.pages) > settings.max_pdf_pages:
        raise ParseError('PDF 页数超过配置限制')
    chunks = []
    text_size = 0
    document = None
    try:
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            has_image = _has_image(page.get('/Resources',{}))
            if len(text.strip()) < 80 and has_image:
                if float(page.mediabox.width)*float(page.mediabox.height)*2.25 > settings.max_ocr_pixels:
                    raise ParseError('PDF扫描页尺寸超过OCR配置上限')
                import pypdfium2 as pdfium
                from .ocr import image_text
                if document is None:
                    document = pdfium.PdfDocument(str(path))
                bitmap = document[index].render(scale=1.5)
                image = bitmap.to_pil()
                try:
                    ocr = image_text(image)
                    text = ocr if len(ocr.strip()) > len(text.strip()) else text
                finally:
                    image.close(); bitmap.close()
            text_size += len(text)
            if text_size > settings.max_document_chars:
                raise ParseError('PDF文本量超过配置上限')
            numbered = '\n'.join(f'[PAGE:{index+1} LINE:{line_number}] {line}'
                for line_number, line in enumerate(text.splitlines(), 1) if line.strip())
            chunks.append(f'[PAGE:{index+1}]\n{numbered}')
    except Exception as exc:
        raise ParseError(f'PDF 解析或 OCR 未完成: {type(exc).__name__}: {str(exc)[:180]}') from exc
    finally:
        if document is not None:
            document.close()
    if not any(part.split('\n',1)[-1].strip() for part in chunks):
        raise ParseError('PDF 未得到可用正文')
    return _bounded_join(chunks)


def _docx_table_rows(table):
    # Match python-docx row.cells, including omitted cells and merged spans.
    # Cache each XML cell once: vertical continuations otherwise repeatedly
    # traverse all prior rows and re-read the root cell's paragraphs.
    from docx.table import _Cell
    resolved = {}

    def values(tc):
        if tc not in resolved:
            if tc.vMerge == 'continue':
                resolved[tc] = values(tc._tc_above)
            else:
                resolved[tc] = [_Cell(tc, table).text.strip()] * tc.grid_span
        return resolved[tc]

    for index, row in enumerate(table.rows, 1):
        yield index, [text for tc in row._tr.tc_lst for text in values(tc)]


def _docx(path: Path) -> str:
    document = Document(str(path))
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    chunks = []
    table_index = 0
    image_index = 0
    seen_images = set()

    def image_chunks(node):
        nonlocal image_index
        for drawing in node.iter():
            if not drawing.tag.endswith('}blip'):
                continue
            relation_id = drawing.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if not relation_id or relation_id in seen_images:
                continue
            seen_images.add(relation_id)
            part = document.part.related_parts.get(relation_id)
            if part is None or not part.content_type.startswith('image/'):
                continue
            try:
                with Image.open(io.BytesIO(part.blob)) as image:
                    if image.width < 80 or image.height < 30:
                        continue
                    if image.width * image.height > settings.max_ocr_pixels:
                        raise ParseError('DOCX内嵌图片尺寸超过OCR配置上限')
                    from .ocr import image_text
                    text = image_text(image)
            except (OSError, ValueError) as exc:
                raise ParseError(f'DOCX内嵌图片无法读取: {type(exc).__name__}') from exc
            image_index += 1
            chunks.extend(f'[IMAGE:{image_index} LINE:{line}] {value}'
                for line, value in enumerate(text.splitlines(), 1) if value.strip())

    for node in document.element.body.iterchildren():
        if node.tag.endswith('}p'):
            text = Paragraph(node,document).text
            if text.strip():
                chunks.append(text)
        elif node.tag.endswith('}tbl'):
            table_index += 1
            rows = _docx_table_rows(Table(node, document))
            chunks.append(_bounded_join(_table_rows(rows, table_index)))
        image_chunks(node)
    return _bounded_join(chunks)


def _xlsx(path: Path) -> str:
    if detect_type(path) == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path),on_demand=True)
        try:
            return _bounded_join(f'[SHEET:{s.name} ROW:{r+1}] '+' | '.join(str(s.cell_value(r,c)) for c in range(s.ncols)) for s in book.sheets() for r in range(s.nrows))
        finally:
            book.release_resources()
    with path.open('rb') as stream:
        book = load_workbook(stream,read_only=True,data_only=True)
        try:
            return _bounded_join(f'[SHEET:{s.title} ROW:{r}] '+' | '.join('' if v is None else str(v) for v in row) for s in book for r,row in enumerate(s.iter_rows(values_only=True),1) if any(v is not None for v in row))
        finally:
            book.close()


def _legacy_document(path: Path) -> str:
    if detect_type(path) == '.rtf':
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(decode_text(path.read_bytes()))
    if detect_type(path) == '.doc' and shutil.which('antiword'):
        process = subprocess.run(['antiword',str(path)],capture_output=True,timeout=45,stdin=subprocess.DEVNULL)
        if process.returncode == 0:
            return decode_text(process.stdout)
    executable = office_tool()
    if not executable:
        raise ParseError('DOC/WPS/RTF 转换工具不可用；需要 LibreOffice 或 antiword')
    with tempfile.TemporaryDirectory(prefix='document-',dir=settings.upload_dir) as directory:
        target = Path(directory)
        process = subprocess.run([executable,'--headless',f'-env:UserInstallation={(target/"profile").as_uri()}',
                                  '--convert-to','docx:Office Open XML Text','--outdir',str(target),str(path)],
                                 capture_output=True,timeout=90,stdin=subprocess.DEVNULL,
                                 env={**os.environ,'SAL_DISABLE_OPENCL':'1','SAL_DISABLESKIA':'1'})
        output = target/(path.stem+'.docx')
        if process.returncode or not output.exists():
            raise ParseError('旧格式文档转换失败')
        return _docx(output)


def _parse_uncached(path: Path) -> str:
    if path.stat().st_size > settings.max_document_bytes:
        raise ParseError('单份文档超过配置上限，请拆分处理')
    kind = detect_type(path)
    if kind in {'.html','.txt','.csv'} and path.stat().st_size > min(settings.max_document_bytes,16*1024**2):
        raise ParseError('单份文本文件过大，请拆分处理')
    if kind in {'.docx','.xlsx'}:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members)>settings.max_archive_files or sum(x.file_size for x in members)>settings.max_document_bytes:
                raise ParseError('Office文档内部累计大小或文件数超限')
            if any(x.filename.lower().endswith('.xml') and x.file_size>32*1024**2 for x in members):
                raise ParseError('Office文档内部XML过大')
    if kind == '.html': text = _html(path)
    elif kind == '.pdf': text = _pdf(path)
    elif kind == '.docx': text = _docx(path)
    elif kind in {'.xlsx','.xls'}: text = _xlsx(path)
    elif kind in {'.doc','.wps','.rtf'}: text = _legacy_document(path)
    elif kind in {'.png','.jpg','.jpeg','.webp'}:
        from .ocr import image_text
        with Image.open(path) as image:
            if image.width*image.height > settings.max_ocr_pixels:
                raise ParseError('图片尺寸超过OCR配置上限')
            text = image_text(image)
    elif kind == '.csv': text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(decode_text(path.read_bytes()))))
    elif kind == '.txt': text = decode_text(path.read_bytes())
    else: raise ParseError('文件需要先解压或格式不受支持')
    if len(text)>settings.max_document_chars:
        raise ParseError('文档文本量超过配置上限')
    if not text.strip():
        raise ParseError('未解析到文本')
    return clean_text(text)


def parse_file(path: Path) -> str:
    if not settings.parse_cache_dir:
        with _PARSE_SLOTS:
            return _parse_uncached(path)
    digest = hashlib.sha256((PARSER_VERSION + str(settings.max_pdf_pages) + str(settings.max_document_chars)
        + str(settings.max_ocr_pixels) + path.suffix.lower()).encode())
    if path.stat().st_size > settings.max_document_bytes:
        raise ParseError('单份文档超过配置上限，请拆分处理')
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024*1024), b''):
            digest.update(block)
    directory = Path(settings.parse_cache_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / (digest.hexdigest() + '.txt')
    if target.is_file():
        text = target.read_text('utf-8')
        if text and len(text) <= settings.max_document_chars:
            return text
    with _PARSE_SLOTS:
        # Another worker can have parsed the same attachment in the meantime.
        if target.is_file():
            return target.read_text('utf-8')
        text = _parse_uncached(path)
    temporary = directory / (digest.hexdigest() + '.' + uuid.uuid4().hex + '.part')
    try:
        temporary.write_text(text, encoding='utf-8')
        temporary.replace(target)
    finally:
        if temporary.exists() and temporary.resolve().is_relative_to(directory):
            temporary.unlink()
    return text


def group_documents(paths: list[Path]) -> list[list[Path]]:
    htmls = sorted([p for p in paths if p.suffix.lower() in {'.html','.htm'}],key=str)
    if not htmls:
        return [[p] for p in sorted(paths,key=str)]
    groups = {p:[p] for p in htmls}
    for attachment in (p for p in paths if p not in htmls):
        candidates = [h for h in htmls if h.stem == attachment.stem or h.stem in attachment.parts]
        if not candidates and len(htmls) == 1:
            candidates = htmls
        if len(candidates) != 1:
            raise ParseError(f'附件无法唯一关联公告: {attachment.name}')
        groups[candidates[0]].append(attachment)
    return list(groups.values())
