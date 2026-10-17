import asyncio
import hashlib
import io
import json

import pytest
from openpyxl import load_workbook

from app.database import one, upsert_notice
from app.evidence import ground_evidence
from app.evaluate import evaluate, score_items
from app.exports import export_records, workbook_bytes
from app.field_text import interpret_field_text
from app.models import ItemExtract, LotExtract, ProjectExtract


@pytest.mark.parametrize('text', [
    '详见投标文件', '见附件', '具体见报价表', '详见附件1《型号清单》', '详见附件3供货范围明细表',
    '满足招标文件要求，详见投标文件“三、分项报价表”', '按招标文件执行',
    '满足本项目磋商文件要求', '符合本次招标文件要求', '按该项目采购文件执行',
])
def test_pure_reference_is_missing_value_with_an_explicit_reason(text):
    observed = interpret_field_text(text, 'brand')
    assert observed.raw == text and observed.value == ''
    assert observed.status == 'reference_only'
    assert ItemExtract(brand=text, model=text).brand == ''


@pytest.mark.parametrize(('text','field','expected'), [
    ('详见附件，品牌为华为','brand','华为'),
    ('品牌：华为，详见附件','brand','华为'),
    ('见投标文件；规格型号：ABC-123','model','ABC-123'),
])
def test_explicit_value_is_preserved_inside_mixed_reference(text,field,expected):
    observed = interpret_field_text(text,field)
    assert observed.value == expected and observed.raw == text
    assert observed.status == 'disclosed_with_reference'
    assert getattr(ItemExtract(**{field:text}),field) == expected


@pytest.mark.parametrize('text', [
    '详见附件 华为', '详见附件，华为', '详见附件，品牌为华为；品牌为联想',
])
def test_ambiguous_mixed_reference_is_not_erased_or_guessed(text):
    assert interpret_field_text(text,'brand').value == text
    assert interpret_field_text(text,'brand').status == 'mixed_reference'


def test_a_label_for_a_different_field_cannot_fill_the_model():
    assert ItemExtract(model='详见附件，品牌为华为').model == '详见附件，品牌为华为'


def source_record():
    record = ProjectExtract(notice_uuid='t20260101_10001',project_name='测试项目',
        lots=[LotExtract(lot_name='包1',items=[ItemExtract(product_name='服务器',source_file='doc-1:a.html',line_id='TABLE:1 ROW:2')])])
    text = '[TABLE:1 ROW:1] 采购标的 | 品牌 | 规格型号\n[TABLE:1 ROW:2] 服务器 | 详见投标文件 | 详见附件'
    ground_evidence(record,text)
    return record


def test_field_reasons_come_from_source_cells_not_model_metadata():
    record = source_record()
    item = record.lots[0].items[0]
    assert item.brand == item.model == ''
    assert item.field_observations['brand']['raw'] == '详见投标文件'
    assert item.field_observations['brand']['source_location'] == 'TABLE:1 ROW:2'
    assert item.field_observations['brand']['status'] == 'reference_only'
    assert item.source_context['text'].endswith('详见投标文件 | 详见附件')
    ground_evidence(record,'[TABLE:1 ROW:2] 服务器')
    assert item.field_observations == {}


def test_ambiguous_source_row_does_not_invent_a_missing_reason():
    record = source_record()
    ground_evidence(record,'[TABLE:1 ROW:2] 服务器\n[TABLE:1 ROW:3] 服务器')
    assert not record.lots[0].items[0].field_observations


def test_ambiguous_reference_is_sent_to_review_instead_of_silent_completion():
    record = source_record()
    ground_evidence(record,'[TABLE:1 ROW:1] 采购标的 | 品牌 | 规格型号\n[TABLE:1 ROW:2] 服务器 | 详见附件，华为 | M1')
    assert record.processing_status == 'partial'
    assert record.lots[0].items[0].field_observations['brand']['status'] == 'mixed_reference'
    assert any('保留原文待核验' in value for value in record.warnings)


def test_field_observations_survive_database_api_and_standard_export(client):
    record = source_record()
    project_id = upsert_notice(record,['a.html'],'reference-fields')[0]
    response = client.get(f'/api/projects/{project_id}',params={'dataset_id':'reference-fields'})
    assert response.status_code == 200
    observation = response.json()['items'][0]['field_observations']['brand']
    assert observation['raw'] == '详见投标文件' and observation['status'] == 'reference_only'
    records,_ = export_records('reference-fields')
    assert records[0].lots[0].items[0].field_observations['brand'] == observation
    book = load_workbook(io.BytesIO(workbook_bytes(records)),data_only=True)
    assert book.sheetnames == ['总表','标的表','投标人表']
    assert book['标的表']['D2'].value == '服务器'
    assert book['标的表']['F2'].value is None and book['标的表']['G2'].value is None
    book.close()


def row(brand=''):
    return {'product_name':'服务器','brand':brand,'source_file':'doc-1:a.html','line_id':'TABLE:1 ROW:2'}


def test_literal_policy_preserves_historical_penalty_and_business_counts_true_negative():
    expected, actual = [row('详见投标文件')], [row()]
    assert score_items(expected,actual)['fields']['brand']['fn'] == 1
    result = score_items(expected,actual,'business-v1')
    assert result['fields']['brand']['tn'] == 1 and result['fields']['brand']['tp'] == 0
    assert result['fields']['brand']['fn'] == 0
    assert result['field_normalizations'][0]['raw'] == '详见投标文件'
    assert expected[0]['brand'] == '详见投标文件'


def test_business_policy_still_penalizes_real_omissions_hallucinations_and_missing_rows():
    assert score_items([row('华为')],[row()],'business-v1')['fields']['brand']['fn'] == 1
    assert score_items([row('详见附件')],[row('某品牌')],'business-v1')['fields']['brand']['fp'] == 1
    missing = score_items([row('详见附件')],[],'business-v1')
    assert missing['missing_rows'] == 1 and missing['fields']['product_name']['fn'] == 1
    assert missing['fields']['brand']['tn'] == 0
    assert score_items([row('华为')],[row('详见附件，品牌为华为')],'business-v1')['fields']['brand']['tp'] == 1


def test_attachment_answer_is_not_excused_as_an_html_reference():
    expected = {**row('华为'),'source_file':'doc-2:报价.pdf','line_id':'PAGE:1 LINE:2'}
    actual = {**expected,'brand':''}
    assert score_items([expected],[actual],'business-v1')['fields']['brand']['fn'] == 1


def test_saved_prediction_evaluation_records_policy_hashes_and_excluded_rows(tmp_path):
    gold = tmp_path/'gold.jsonl'
    predictions = tmp_path/'predictions.jsonl'
    gold.write_text(json.dumps({'notice_uuid':'n','items':[row('详见附件')]},ensure_ascii=False)+'\n',encoding='utf-8')
    excluded = {**row('华为'),'line_id':'TABLE:2 ROW:2'}
    predictions.write_text(json.dumps({'notice_uuid':'n','items':[row(),excluded]},ensure_ascii=False)+'\n',encoding='utf-8')
    original = gold.read_bytes()
    report = asyncio.run(evaluate(gold,predictions,labelled_tables_only=True,field_policy='business-v1'))
    assert report['field_policy'] == 'business-v1' and report['official_score'] is False
    assert report['micro']['fp'] == report['micro']['fn'] == 0
    assert report['diagnostics'][0]['unlabelled_table_rows'] == 1
    assert report['diagnostics'][0]['field_normalizations'][0]['raw'] == '详见附件'
    assert report['model_calls'] == 0
    assert report['provenance']['gold_sha256'] == hashlib.sha256(original).hexdigest()
    assert report['provenance']['predictions_sha256'] == hashlib.sha256(predictions.read_bytes()).hexdigest()
    assert gold.read_bytes() == original
    with pytest.raises(ValueError,match='Unknown field policy'):
        asyncio.run(evaluate(gold,predictions,field_policy='unknown'))
