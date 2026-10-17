import asyncio
import json
from collections import Counter
from dataclasses import replace

import pytest
import httpx
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app import extractor, parsers
from app.config import settings
from app.models import ProjectExtract


def resume_fixture(tmp_path, monkeypatch):
    configuration = replace(settings, deepseek_api_key='fixture-only', model_network_enabled=False,
                            model_cache_dir='', document_cache_dir=str(tmp_path / 'pieces'))
    files = [tmp_path / name for name in ('notice.html', 'quote.txt', 'other.txt')]
    for path, text in zip(files, ('项目原文', '报价原文', '其他原文')):
        path.write_text(text, encoding='utf-8')
    monkeypatch.setattr(extractor, 'parse_file', lambda path: path.read_text('utf-8'))
    calls = Counter()

    async def call(self, text, names):
        calls[names[0]] += 1
        return ProjectExtract(project_name='已披露采购项目', purchaser='采购单位')

    monkeypatch.setattr(extractor.DeepSeekExtractor, '_call', call)
    return configuration, files, calls, call


def run(configuration, files):
    model = extractor.DeepSeekExtractor(configuration=configuration)
    return asyncio.run(model.extract(files)), model


def without_usage(result):
    return result.model_dump(exclude={'model_usage'})


def test_resume_only_retries_unsuccessful_attachment(tmp_path, monkeypatch):
    configuration, files, calls, original = resume_fixture(tmp_path, monkeypatch)

    async def fail_once(self, text, names):
        if names[0].endswith('quote.txt') and not calls[names[0]]:
            calls[names[0]] += 1
            raise RuntimeError('暂时不可用')
        return await original(self, text, names)

    monkeypatch.setattr(extractor.DeepSeekExtractor, '_call', fail_once)
    first, _ = run(configuration, files)
    assert first.processing_status == 'partial'
    assert first.model_usage['document_cache_writes'] == 2
    second, model = run(configuration, files)
    assert second.processing_status == 'complete'
    assert second.model_usage['document_cache_hits'] == 2
    assert model.document_cache_writes == 1
    assert calls == {'doc-1:notice.html': 1, 'doc-2:quote.txt': 2, 'doc-3:other.txt': 1}
    clean, _ = run(replace(configuration, document_cache_dir=''), files)
    assert without_usage(second) == without_usage(clean)


def test_checkpoint_reads_never_create_http_client(tmp_path, monkeypatch):
    configuration, files, calls, _ = resume_fixture(tmp_path, monkeypatch)
    first, _ = run(configuration, files)
    monkeypatch.setattr(extractor.DeepSeekExtractor, '_call', lambda *args: pytest.fail('No model call expected'))
    monkeypatch.setattr(extractor.httpx, 'AsyncClient', lambda *args, **kwargs: pytest.fail('No HTTP client expected'))
    second, model = run(configuration, files)
    assert without_usage(first) == without_usage(second)
    assert model.calls == second.model_usage['requests'] == 0
    assert second.model_usage['document_cache_hits'] == 3


@pytest.mark.parametrize('changed', ['source', 'prompt', 'model', 'tokens', 'pipeline', 'announcement'])
def test_changed_inputs_cannot_reuse_grounded_results(tmp_path, monkeypatch, changed):
    configuration, files, calls, _ = resume_fixture(tmp_path, monkeypatch)
    run(configuration, files)
    if changed == 'source':
        files[1].write_text('新的报价原文', encoding='utf-8')
    elif changed == 'announcement':
        files[0].write_text('新的公告上下文', encoding='utf-8')
    elif changed == 'prompt':
        monkeypatch.setattr(extractor, 'system_prompt', lambda strategy: '新规则')
    elif changed == 'pipeline':
        monkeypatch.setattr(extractor, 'PIPELINE_VERSION', 'changed-rules')
    elif changed == 'model':
        configuration = replace(configuration, deepseek_model='another-model')
    elif changed == 'tokens':
        configuration = replace(configuration, model_max_tokens=2048)
    result, _ = run(configuration, files)
    assert result.processing_status == 'complete'
    assert calls['doc-2:quote.txt'] == (1 if changed == 'announcement' else 2)
    if changed == 'source':
        assert result.model_usage['document_cache_hits'] == 2
    elif changed == 'announcement':
        # Attachment context is the unchanged grounded project, not a raw HTML prefix.
        assert calls['doc-1:notice.html'] == 2
    else:
        assert result.model_usage['document_cache_hits'] == 0


def test_partial_piece_is_not_frozen_as_success(tmp_path, monkeypatch):
    configuration, files, calls, original = resume_fixture(tmp_path, monkeypatch)

    async def partial(self, text, names):
        result = await original(self, text, names)
        result.processing_status = 'partial'
        return result

    monkeypatch.setattr(extractor.DeepSeekExtractor, '_call', partial)
    first, _ = run(configuration, files)
    second, model = run(configuration, files)
    assert first.processing_status == second.processing_status == 'partial'
    assert model.document_cache_hits == model.document_cache_writes == 0
    assert all(count == 2 for count in calls.values())


@pytest.mark.parametrize('damage', ['content', 'metadata', 'invalid-json'])
def test_corrupt_checkpoint_falls_back_to_original_extraction(tmp_path, monkeypatch, damage):
    configuration, files, calls, _ = resume_fixture(tmp_path, monkeypatch)
    first, _ = run(configuration, files)
    for path in (tmp_path / 'pieces').glob('*.json'):
        record = json.loads(path.read_text('utf-8'))
        if damage == 'content':
            record['pieces'][0]['purchaser'] = '被改动的缓存'
        elif damage == 'metadata':
            record['metadata'] = None
        path.write_text('{' if damage == 'invalid-json' else json.dumps(record), encoding='utf-8')
    second, _ = run(configuration, files)
    assert without_usage(first) == without_usage(second)
    assert all(count == 2 for count in calls.values())


def test_unwritable_cache_does_not_break_extraction(tmp_path, monkeypatch):
    configuration, files, calls, _ = resume_fixture(tmp_path, monkeypatch)
    obstruction = tmp_path / 'not-a-directory'
    obstruction.write_text('file', encoding='utf-8')
    result, model = run(replace(configuration, document_cache_dir=str(obstruction)), files)
    assert result.processing_status == 'complete'
    assert model.document_cache_writes == 0
    assert len(calls) == 3


def test_resume_retains_successful_chunks_inside_one_attachment(tmp_path, monkeypatch):
    configuration, files, calls, original = resume_fixture(tmp_path, monkeypatch)
    files[0].write_text('\n'.join(f'{index:03}分项原文' for index in range(80)), encoding='utf-8')
    configuration = replace(configuration, max_llm_chars=512)

    async def fail_second_chunk(self, text, names):
        if calls[names[0]] == 1:
            calls[names[0]] += 1
            raise RuntimeError('第二个文本块暂时失败')
        return await original(self, text, names)

    monkeypatch.setattr(extractor.DeepSeekExtractor, '_call', fail_second_chunk)
    first, _ = run(configuration, files[:1])
    assert first.processing_status == 'partial'
    before = calls['doc-1:notice.html']
    second, _ = run(configuration, files[:1])
    assert second.processing_status == 'complete'
    assert calls['doc-1:notice.html'] == before + 1
    assert second.model_usage['document_cache_hits'] == first.model_usage['document_cache_writes'] > 1


def test_changed_rag_knowledge_invalidates_checkpoints(tmp_path, monkeypatch):
    from app import retrieval
    configuration, files, calls, _ = resume_fixture(tmp_path, monkeypatch)
    knowledge = tmp_path / 'knowledge.json'
    knowledge.write_text('first corpus', encoding='utf-8')
    monkeypatch.setattr(retrieval, 'KNOWLEDGE', knowledge)
    configuration = replace(configuration, extraction_strategy='rag')
    run(configuration, files)
    cached, _ = run(configuration, files)
    assert cached.model_usage['document_cache_hits'] == 3
    knowledge.write_text('changed corpus', encoding='utf-8')
    changed, _ = run(configuration, files)
    assert changed.model_usage['document_cache_hits'] == 0
    assert all(value == 2 for value in calls.values())


def test_more_inflight_notices_cannot_exceed_model_request_limit(tmp_path):
    configuration = replace(settings, deepseek_api_key='fixture-only', model_network_enabled=True,
                            model_concurrency=2, model_require_free=False,
                            model_cache_dir='', document_cache_dir='')
    active = peak = 0

    async def handler(request):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        try:
            await asyncio.sleep(0.01)
            return httpx.Response(200, json={'choices': [{'finish_reason': 'stop',
                'message': {'content': '{"project_name":"采购项目"}'}}]})
        finally:
            active -= 1

    async def check():
        model = extractor.DeepSeekExtractor(configuration=configuration)
        model.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        try:
            records = await asyncio.gather(*(model._call(str(index), ['notice.html']) for index in range(8)))
            assert len(records) == model.calls == 8
            assert peak == 2 and active == 0
        finally:
            await model.close()

    asyncio.run(check())


def test_docx_merged_and_omitted_cells_match_original_parser():
    document = Document()
    table = document.add_table(rows=25, cols=4)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = f'行{row_index}列{column_index}'
    table.cell(0, 0).merge(table.cell(20, 0))
    table.cell(21, 1).merge(table.cell(21, 3))
    row = table.rows[24]
    row._tr.remove(row._tr.tc_lst[0])
    grid_before = OxmlElement('w:gridBefore')
    grid_before.set(qn('w:val'), '1')
    row._tr.get_or_add_trPr().append(grid_before)
    expected = [(index, [cell.text.strip() for cell in row.cells])
                for index, row in enumerate(table.rows, 1)]
    assert list(parsers._docx_table_rows(table)) == expected
