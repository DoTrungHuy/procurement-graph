"""Exercise a clean source bundle using existing Python dependencies; model HTTP forbidden."""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import subprocess
import sys
import tempfile
import zipfile

ROOT=Path(__file__).resolve().parents[1]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--bundle',type=Path,default=ROOT/'deliverables/procurement-release.zip')
    args=parser.parse_args()
    staging=ROOT/'.cache/release-check'
    staging.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='source-',dir=staging) as temporary:
        scratch=Path(temporary).resolve()
        with zipfile.ZipFile(args.bundle) as archive:
            for member in archive.infolist():
                relative=PurePosixPath(member.filename)
                if relative.is_absolute() or '..' in relative.parts or any(':' in part for part in relative.parts):
                    raise ValueError('Invalid bundle path')
                target=scratch.joinpath(*relative.parts)
                if not target.resolve().is_relative_to(scratch):
                    raise ValueError('Bundle path escapes staging')
                if member.is_dir():
                    target.mkdir(parents=True,exist_ok=True)
                else:
                    target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(archive.read(member))
        source=scratch/'procurement-project'
        manifest=json.loads((source/'release-manifest.json').read_text('utf-8'))
        for name,metadata in manifest['files'].items():
            if hashlib.sha256((source/name).read_bytes()).hexdigest()!=metadata['sha256']:
                raise ValueError('Bundle checksum mismatch: '+name)
        assert not (source/'.env').exists()
        assert not (source/'demo/procurement/data/entity_mining.db').exists()
        assert not (source/'data/official-benchmark/html').exists()
        env={**os.environ,'PYTHONPATH':str(source/'demo/procurement/backend'),'PYTHONIOENCODING':'utf-8',
             'RELEASE_SOURCE_ROOT':str(source),
             'RELEASE_EXPECTED_VERSION':manifest['application_version'],
             'DATABASE_PATH':str(scratch/'fresh.db'),'UPLOAD_DIR':str(scratch/'uploads'),
             'MODEL_SETTINGS_PATH':str(scratch/'model-settings.local.json'),
             'PARSE_CACHE_DIR':str(scratch/'parsed'),'MODEL_RESPONSE_CACHE_DIR':str(source/'demo/procurement/data/demo/cache'),
             'OFFICIAL_DATA_DIR':str(scratch/'absent-official-input'),'REFERENCE_SEED':'true','DEMO_SEED':'false',
             'DEEPSEEK_API_KEY':'','DEEPSEEK_BASE_URL':'https://api.deepseek.com','DEEPSEEK_MODEL':'deepseek-flash',
             'MODEL_MAX_TOKENS':'16384','MODEL_THINKING':'false','EXTRACTION_STRATEGY':'schema','MODEL_NETWORK_ENABLED':'false',
             'APP_API_TOKEN':'','NEO4J_URI':'','RELATION_BACKEND':'sqlite'}
        program='''import io,json,httpx,os,runpy
runpy.run_path(os.environ['RELEASE_SOURCE_ROOT']+'/scripts/prepare_offline_demo.py',run_name='__main__')
from fastapi.testclient import TestClient
from openpyxl import load_workbook
async def forbidden(*args,**kwargs): raise RuntimeError('Model HTTP forbidden in release verification')
httpx.AsyncClient.request=forbidden
from app.main import app
with TestClient(app) as client:
    health=client.get('/api/health').json()
    assert health['version']==os.environ['RELEASE_EXPECTED_VERSION']
    assert not health['deepseek_configured'] and not health['model_network_enabled']
    assert not health['official_data_available']
    stats=client.get('/api/stats',params={'dataset_id':'official-reference'}).json()
    assert stats['notices']==20 and stats['projects']==26,stats
    created=client.post('/api/samples/replay',json={'dataset_id':'release-check'}).json()
    job=client.get('/api/jobs/'+created['id']).json()
    assert job['finished'] and job['failed']==20 and job['metrics']['model_calls']==0,job
    assert job['metrics']['response_cache_hits']==0,job
    assert all('未命中当前规则' in error for error in job['errors']),job
    assert client.get('/api/stats',params={'dataset_id':'release-check'}).json()['notices']==0
    exported=client.get('/api/export/standard.xlsx',params={'dataset_id':'official-reference'})
    assert exported.headers['content-disposition']=='attachment; filename=procurement-tables.xlsx'
    workbook=exported.content
    names=load_workbook(io.BytesIO(workbook),read_only=True).sheetnames
    assert names==['总表','标的表','投标人表'],names
    print(json.dumps({'passed':True,'application_version':health['version'],'notices':20,'packages':26,'cache_hits':job['metrics']['response_cache_hits'],
                      'model_requests':0,'workbook_bytes':len(workbook),'database':'fresh','official_input':'absent',
                      'data_source':'bundled historical snapshot, not current-rule extraction',
                      'stale_replay_rejected':True}))
'''
        completed=subprocess.run([sys.executable,'-X','utf8','-c',program],cwd=ROOT,env=env,capture_output=True,text=True,encoding='utf-8')
        if completed.returncode:
            print(completed.stdout);print(completed.stderr,file=sys.stderr);raise SystemExit(completed.returncode)
        result=json.loads(completed.stdout.strip().splitlines()[-1])
    result.update(checked_at=datetime.now().astimezone().isoformat(),temporary_source_removed=True,
                  scope='Clean extracted source; existing Python dependencies; not a fresh dependency installation or official platform test')
    output=ROOT/'.cache/maintenance/release-verification.json'
    output.parent.mkdir(parents=True,exist_ok=True);output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    main()
